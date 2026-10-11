from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

SOURCE = Path(r"C:\Users\kaziNAHIAN\Downloads\notes.docx")
OUTPUT = Path(r"C:\user\juli\thesis\original_file\all_codes\notes_remaining_refs_updated.docx")


def set_font(run, bold=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(12)
    run.bold = bold


def add(doc, text, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_font(r, bold)


def add_section(doc, marker, entries):
    if any(p.text.strip() == marker for p in doc.paragraphs):
        return
    add(doc, marker, True)
    for text, bold in entries:
        add(doc, text, bold)


doc = Document(SOURCE)

add_section(doc, "Reference [46]", [
    ("The current paper concerns histogram-equalization-based tumor detection and does not directly support the statement about scaling pixel values before neural-network training. The replacement is a peer-reviewed systematic review devoted to MRI intensity standardization in glioma research.", False),
    ("Delete current reference", True),
    ("M. Ramamoorthy, S. Qamar, R. Manikandan, N. Z. Jhanjhi, M. Masud, and M. A. AlZain, \"Earlier detection of brain tumor by pre-processing based on histogram equalization with neural network,\" Healthcare, vol. 10, no. 7, 2022.", False),
    ("Replace it with", True),
    ("K. Fatania, F. Mohamud, A. Clark, M. Nix, S. C. Short, J. O'Connor, A. F. Scarsbrook, and S. Currie, \"Intensity standardization of MRI prior to radiomic feature extraction for artificial intelligence research in glioma: a systematic review,\" European Radiology, vol. 32, no. 10, pp. 7014-7025, 2022, doi: 10.1007/s00330-022-08807-2.", False),
    ("Delete this sentence", True),
    ("This normalization ensures that all images are represented on a consistent intensity scale and prevents large pixel values from dominating gradient updates [46].", False),
    ("Replace it with", True),
    ("Each RGB channel was rescaled from the original 0-255 range to 0-1 by dividing its pixel values by 255. This intensity normalization places all input images on a common bounded scale before model training [46].", False),
])

add_section(doc, "Reference [49]", [
    ("The citation is attached to the mathematical definition of Otsu thresholding. The most appropriate source is therefore Otsu's original peer-reviewed article rather than a later application paper on morphological operations and watershed segmentation.", False),
    ("Delete current reference", True),
    ("D. Mishra et al., \"Brain tumor detection using morphological operations and watershed segmentation,\" Multimedia Tools and Applications, vol. 79, pp. 19007-19026, 2020.", False),
    ("Replace it with", True),
    ("N. Otsu, \"A threshold selection method from gray-level histograms,\" IEEE Transactions on Systems, Man, and Cybernetics, vol. 9, no. 1, pp. 62-66, 1979, doi: 10.1109/TSMC.1979.4310076.", False),
    ("Sentence beside Equation (4)", True),
    ("Keep the existing explanation of Otsu's within-class variance and retain citation [49] after that explanation. No substantive sentence rewrite is required because the equation and definitions match the original Otsu method.", False),
])

add_section(doc, "References [59]-[62]", [
    ("These are relevant recent journal studies, but their present citations are attached to broad limitations and future-work statements that the papers do not directly establish. Keep the papers, correct the incomplete bibliography details below, cite them in the literature review where their actual methods and results are discussed, and remove the misplaced citation numbers from the limitations section.", False),
    ("Correct bibliography entries", True),
    ("[59] T. Banerjee, P. Chhabra, M. Kumar, A. Kumar, K. Abhishek, and M. A. Shah, \"Pyramidal attention-based T network for brain tumor classification: a comprehensive analysis of transfer learning approaches for clinically reliable and reliable AI hybrid approaches,\" Scientific Reports, vol. 15, Art. no. 28669, 2025, doi: 10.1038/s41598-025-11574-x.", False),
    ("[60] U. Ali, Sajid, R. Saleem, N. Imran, S. A. O. Bahaj, T. Mahmood, and N. Ayesha, \"A lightweight deep learning framework for real time brain tumor detection and characterization using MR images,\" Discover Artificial Intelligence, vol. 6, Art. no. 341, 2026, doi: 10.1007/s44163-026-01046-0.", False),
    ("[61] V. Mehrdad, R. Talebzadeh, and N. Fazaeli, \"Brain tumor classification using optimized ResNet50 with dynamic precision optimization for enhanced speed and diagnostic accuracy,\" Scientific Reports, vol. 16, Art. no. 9263, 2026, doi: 10.1038/s41598-026-39926-1.", False),
    ("[62] M. A. A. Sheikh, A. Dey, and A. Kole, \"A deep transfer learning and ensemble learning framework for brain tumor detection and classification in MRI images,\" Discover Artificial Intelligence, vol. 6, Art. no. 859, 2026, doi: 10.1007/s44163-026-01080-y.", False),
    ("Add this paragraph to the literature review", True),
    ("Recent studies have examined attention-based, lightweight, efficiency-oriented, and ensemble approaches for brain tumor classification. Banerjee et al. [59] proposed PABT-Net, which combines hierarchical pyramidal attention with bi-partitioned feature extraction and reported 99.12% overall accuracy and 98.77% mean accuracy under nine-fold cross-validation across three benchmark datasets. Ali et al. [60] introduced FCDS-DeepVision, a lightweight CNN trained from scratch with MRI-specific normalization, noise reduction, contrast enhancement, and foreground localization. Mehrdad et al. [61] developed an optimized ResNet50 framework with dynamic precision management and a ResNet-Random Forest component, reporting 99.69% accuracy on a four-class dataset of 7,023 images and evaluating generalization using five-fold cross-validation. Sheikh et al. [62] combined transfer learning, accuracy-weighted ensemble voting, and Grad-CAM, achieving 99.28% accuracy on the BT-7023 dataset and 98.80% accuracy in cross-dataset evaluation using an external Figshare dataset.", False),
    ("Correct the citations in the limitations section", True),
    ("1. In the sentence ending 'generalizability of the proposed framework [59], [60],' delete only '[59], [60]'. The statement is your study's limitation and does not require these model papers as evidence.", False),
    ("2. In the sentence ending 'clinically verified tumor regions [61],' delete only '[61]'. This is a proposed direction for your own work.", False),
    ("3. In the sentence ending 'other clinically relevant imaging modalities where appropriate [62],' delete only '[62]'.", False),
    ("4. In the sentence ending 'resource-constrained clinical decision-support applications [60], [61],' delete only '[60], [61]'. Keep the future-work sentence itself.", False),
])

doc.save(OUTPUT)
print(OUTPUT)
