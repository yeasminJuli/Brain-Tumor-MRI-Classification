from pathlib import Path

from docx import Document
from docx.shared import Pt


SOURCE = Path(r"C:\Users\kaziNAHIAN\Downloads\notes.docx")
OUTPUT = Path(r"C:\user\juli\thesis\original_file\all_codes\notes_reference_5_temp.docx")

doc = Document(SOURCE)
marker = "Reference [5]"

if not any(p.text.strip() == marker for p in doc.paragraphs):
    def add(text, style="Normal (Web)", bold=False):
        paragraph = doc.add_paragraph(style=style)
        run = paragraph.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        run.bold = bold
        return paragraph

    add(marker, style="Heading 2")
    add(
        "The current reference is a technical study on automated superpixel-based "
        "tumor segmentation. It does not directly support the manuscript's description "
        "of how neuroimaging is interpreted in clinical practice. Replace it with the "
        "following recent evidence-based imaging guideline published in 2025."
    )
    add("Replace the bibliography entry [5]", style="Heading 3")
    add("Delete the current reference [5]:", bold=True)
    add(
        '[5] Z. U. Rehman, S. S. Naqvi, T. M. Khan, M. A. Khan, and T. Bashir, '
        '"Fully automated multi-parametric brain tumour segmentation using superpixel '
        'based classification," Expert Systems with Applications, vol. 118, '
        'pp. 598-613, 2019.'
    )
    add("Replace it with:", bold=True)
    add(
        '[5] J. Ivanidze et al., "ACR Appropriateness Criteria Brain Tumors," '
        'Journal of the American College of Radiology, vol. 22, no. 5S, '
        'pp. S108-S135, 2025, doi: 10.1016/j.jacr.2025.02.036.'
    )
    add("Delete this sentence from the Introduction:", bold=True)
    add(
        "In clinical practice, radiological images are interpreted by physicians or "
        "radiologists to identify tumor type and guide patient management [5]."
    )
    add("Replace it with:", bold=True)
    add(
        "In clinical practice, neuroimaging is interpreted by radiologists and other "
        "treating physicians to characterize suspected brain tumors and support "
        "diagnostic, pretreatment, and follow-up decisions [5]."
    )

doc.save(OUTPUT)
print(OUTPUT)
