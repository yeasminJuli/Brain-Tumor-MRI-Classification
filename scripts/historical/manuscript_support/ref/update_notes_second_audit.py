from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

SOURCE = Path(r"C:\Users\kaziNAHIAN\Downloads\notes.docx")
OUTPUT = Path(r"C:\user\juli\thesis\original_file\all_codes\notes_second_audit_updated.docx")


def add(doc, text, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    r.font.size = Pt(12)
    r.bold = bold


def section(doc, marker, entries):
    if any(p.text.strip() == marker for p in doc.paragraphs):
        return
    add(doc, marker, True)
    for text, bold in entries:
        add(doc, text, bold)


doc = Document(SOURCE)

section(doc, "Reference [14]", [
    ("The current entry shortens the published title and gives the wrong issue number. The study itself is appropriate, but its bibliographic information should be corrected to match the publisher's record.", False),
    ("Delete current reference", True),
    ("S. Natha et al., \"Ensemble deep learning for brain tumour detection,\" Applied Sciences, vol. 14, no. 6, p. 2210, 2024.", False),
    ("Replace it with", True),
    ("S. Natha, U. Laila, I. A. Gashim, K. Mahboob, M. N. Saeed, and K. M. Noaman, \"Automated brain tumor identification in biomedical radiology images: a multi-model ensemble deep learning approach,\" Applied Sciences, vol. 14, no. 5, Art. no. 2210, 2024, doi: 10.3390/app14052210.", False),
])

section(doc, "Reference [35]", [
    ("The DOI, title, journal, and year are valid, but the Scientific Reports article number is missing. Add article number 19866.", False),
    ("Replace the bibliography entry with", True),
    ("[35] B. D. Shivahare, S. K. Subramaniam, Dafik et al., \"An advanced hybrid deep learning framework for high-precision brain tumor detection and classification in MRI scans,\" Scientific Reports, vol. 16, Art. no. 19866, 2026, doi: 10.1038/s41598-026-50194-x.", False),
])

section(doc, "Reference [57]", [
    ("The number 21 is the article number, not the journal issue number. The DOI should also be included.", False),
    ("Delete current reference", True),
    ("S. Srinivasan, D. Francis, S. K. Mathivanan, H. Rajadurai, B. D. Shivahare, and M. A. Shah, \"A hybrid deep CNN model for brain tumor image multi-classification,\" BMC Medical Imaging, vol. 24, no. 21, 2024.", False),
    ("Replace it with", True),
    ("S. Srinivasan, D. Francis, S. K. Mathivanan, H. Rajadurai, B. D. Shivahare, and M. A. Shah, \"A hybrid deep CNN model for brain tumor image multi-classification,\" BMC Medical Imaging, vol. 24, Art. no. 21, 2024, doi: 10.1186/s12880-024-01195-7.", False),
])

section(doc, "Second-Pass Reference Audit Status", [
    ("The numbered corrections recorded in this notes document are supported by publisher or indexed journal records. They still need to be applied to the manuscript before the manuscript reference list can be considered corrected.", False),
    ("Items requiring an editorial decision", True),
    ("References [8]-[10] are clinical webpages rather than journal articles. If the final bibliography must contain journal articles only, these should be replaced with peer-reviewed tumor-specific sources.", False),
    ("References [42]-[45] describe the datasets used in the study. They are not journal articles, but they should normally remain because the data sources must be cited transparently.", False),
    ("References [50]-[56] are the original papers for attention, Grad-CAM++, LIME, VGG16, ResNet, DenseNet, and EfficientNet, and several are conference publications. These are canonical primary sources. Removing them solely because they are conference papers could weaken methodological attribution; retain them unless the journal or supervisor explicitly requires a journal-only bibliography.", False),
    ("After all noted replacements are applied, perform one final citation-number and IEEE-style formatting check on the revised manuscript because changing bibliography entries alone does not update the surrounding prose automatically.", False),
])

doc.save(OUTPUT)
print(OUTPUT)
