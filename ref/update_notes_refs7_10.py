from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

SOURCE = Path(r"C:\Users\kaziNAHIAN\Downloads\notes.docx")
OUTPUT = Path(r"C:\user\juli\thesis\original_file\all_codes\notes_refs7_10_updated.docx")


def add(doc, text, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    r.font.size = Pt(12)
    r.bold = bold


def section(doc, marker, entries):
    if any(p.text.strip() == marker for p in doc.paragraphs):
        return
    add(doc, marker, True)
    for text, bold in entries:
        add(doc, text, bold)


doc = Document(SOURCE)

section(doc, "References [7]-[10] Journal-Article Replacements", [
    ("References [7]-[10] are presently a WHO classification document and three clinical webpages. Because the final bibliography is required to contain journal articles, replace them with the peer-reviewed sources below. The replacement sentence is also narrowed so that every cited article directly supports the associated statement.", False),
    ("Reference [7]", True),
    ("Delete: World Health Organization. (2021) WHO classification of tumours of the central nervous system. WHO CNS Tumor Classification.", False),
    ("Replace with: D. N. Louis et al., \"The 2021 WHO Classification of Tumors of the Central Nervous System: a summary,\" Neuro-Oncology, vol. 23, no. 8, pp. 1231-1251, 2021, doi: 10.1093/neuonc/noab106.", False),
    ("Reference [8]", True),
    ("Delete: Johns Hopkins Medicine, \"Gliomas,\" 2026, available: https://www.hopkinsmedicine.org/health/conditions-and-diseases/gliomas.", False),
    ("Replace with: M. Weller et al., \"Glioma,\" Nature Reviews Disease Primers, vol. 10, Art. no. 33, 2024, doi: 10.1038/s41572-024-00516-y.", False),
    ("Reference [9]", True),
    ("Delete: Johns Hopkins Medicine, \"Meningioma,\" 2026, available: https://www.hopkinsmedicine.org/health/conditions-and-diseases/meningioma.", False),
    ("Replace with: J. Z. Wang et al., \"Meningioma: International Consortium on Meningiomas consensus review on scientific advances and treatment paradigms for clinicians, researchers, and patients,\" Neuro-Oncology, vol. 26, no. 10, pp. 1742-1780, 2024, doi: 10.1093/neuonc/noae082.", False),
    ("Reference [10]", True),
    ("Delete: Mayo Clinic, \"Pituitary tumors-symptoms and causes,\" 2026, available: https://www.mayoclinic.org/diseases-conditions/pituitary-tumors.", False),
    ("Replace with: F. Roncaroli, C. A. Donofrio, L. Walker, R. Laitt, C. Villa, and W. Majeed, \"Update on the classification and diagnostic approach of pituitary neuroendocrine tumours,\" Diagnostic Histopathology, vol. 30, no. 12, pp. 668-679, 2024, doi: 10.1016/j.mpdhp.2024.10.001.", False),
    ("Replace the associated sentences in the Introduction", True),
    ("Delete: Brain tumors are commonly categorized by histological and molecular characteristics, anatomical origin, and tumor grade. The World Health Organization grading framework classifies central nervous system tumors according to features associated with biological behavior and clinical aggressiveness. Tumor grades range from lower-grade lesions with relatively slower progression to higher-grade tumors characterized by greater cellular abnormality, proliferation rate, and invasiveness, with higher-grade tumors generally associated with more aggressive behavior [7]. Among commonly studied tumor categories, gliomas originate from glial cells, meningiomas arise from the meninges, and pituitary tumors occur within or near the pituitary gland [8]-[10].", False),
    ("Replace with: The 2021 World Health Organization classification integrates histological and molecular characteristics in the diagnosis and grading of central nervous system tumors [7]. Among the commonly studied tumor groups, gliomas, meningiomas, and pituitary neuroendocrine tumors have distinct cellular origins, biological characteristics, and contemporary classification frameworks [8]-[10].", False),
    ("Why this wording is safer", True),
    ("The revised wording avoids treating all tumors as if they follow one simple anatomical definition. It reflects the modern molecular classification of gliomas, the current consensus understanding of meningiomas, and the updated PitNET terminology used for anterior pituitary tumors.", False),
])

doc.save(OUTPUT)
print(OUTPUT)
