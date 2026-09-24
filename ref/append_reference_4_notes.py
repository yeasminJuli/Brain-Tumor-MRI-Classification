from pathlib import Path

from docx import Document
from docx.shared import Pt


SOURCE = Path(r"C:\Users\kaziNAHIAN\Downloads\notes.docx")
OUTPUT = Path(r"C:\user\juli\thesis\original_file\all_codes\notes_reference_4_temp.docx")

doc = Document(SOURCE)

marker = "Reference [4]"
if not any(p.text.strip() == marker for p in doc.paragraphs):
    def add(text, style="Normal (Web)", bold=False):
        paragraph = doc.add_paragraph(style=style)
        run = paragraph.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        run.bold = bold
        return paragraph

    add(marker, style="Heading 2")
    add(
        "The current reference does not directly support the manuscript's statement "
        "about differences in brain-tumor incidence across populations. Replace it "
        "with the following recent global epidemiological study published in 2025."
    )
    add("Replace the bibliography entry [4]", style="Heading 3")
    add("Delete the current reference [4]:", bold=True)
    add(
        '[4] S. L. Fernandes, U. J. Tanik, V. Rajinikanth, and K. A. Karthik, '
        '"A reliable framework for accurate brain image examination and treatment '
        'planning based on early diagnosis support for clinicians," Neural Computing '
        'and Applications, vol. 32, no. 20, pp. 15897-15908, 2020.'
    )
    add("Replace it with:", bold=True)
    add(
        '[4] A. Miranda-Filho, A. Znaor, C. Sunguc, M. Zahwe, R. Marcos-Gragera, '
        'J. D. Figueroa, and F. Bray, "Cancers of the brain and central nervous '
        'system: Global patterns and trends in incidence," Journal of Neuro-Oncology, '
        'vol. 172, no. 3, pp. 567-578, 2025, '
        'doi: 10.1007/s11060-025-04944-y.'
    )
    add("Delete this sentence from the Introduction:", bold=True)
    add(
        "Brain tumors remain associated with substantial mortality and long-term "
        "disability, and their incidence varies across populations and reporting "
        "systems [4]."
    )
    add("Replace it with:", bold=True)
    add(
        "Brain and central nervous system cancers continue to impose a substantial "
        "global burden, with incidence rates varying considerably across countries "
        "and population groups [4]."
    )

doc.save(OUTPUT)
print(OUTPUT)
