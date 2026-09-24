import sys
from docx import Document

sys.stdout.reconfigure(encoding="utf-8")
path = r"C:\Users\kaziNAHIAN\Downloads\notes.docx"
doc = Document(path)
markers = [
    "Reference [6]", "Reference [14]", "Reference [35]", "Reference [46]",
    "Reference [49]", "Reference [57]", "References [59]-[62]",
    "Second-Pass Reference Audit Status",
    "References [7]-[10] Journal-Article Replacements",
]
for marker in markers:
    matches = [p for p in doc.paragraphs if p.text.strip() == marker]
    print(f"{marker}: {len(matches)}")

start = next(i for i,p in enumerate(doc.paragraphs) if p.text.strip() == "Reference [6]")
bad = []
for i, p in enumerate(doc.paragraphs[start:], start):
    for r in p.runs:
        if r.text.strip():
            name = r.font.name
            size = r.font.size.pt if r.font.size else None
            if name != "Times New Roman" or size != 12:
                bad.append((i, r.text[:50], name, size))
print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)}")
print(f"font_issues_after_reference_6={len(bad)}")
for item in bad[:10]:
    print(item)
