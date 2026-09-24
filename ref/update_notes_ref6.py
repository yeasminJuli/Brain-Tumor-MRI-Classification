from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

SOURCE = Path(r"C:\Users\kaziNAHIAN\Downloads\notes.docx")
OUTPUT = Path(r"C:\user\juli\thesis\original_file\all_codes\notes_ref6_updated.docx")
MARKER = "Reference [6]"


def set_font(run, *, bold=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(12)
    run.bold = bold


def add_text(doc, text, *, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_font(r, bold=bold)
    return p


doc = Document(SOURCE)
if not any(p.text.strip() == MARKER for p in doc.paragraphs):
    add_text(doc, MARKER, bold=True)
    add_text(
        doc,
        "The current source concerns texture-based machine-learning localization and does not support the statement about agreement among readers. A recent multi-radiologist study directly evaluates inter-rater agreement for MRI features in adult glioma cases, so it is a better match for this claim.",
    )
    add_text(doc, "Delete current reference", bold=True)
    add_text(
        doc,
        "Z. U. Rehman, M. S. Zia, G. R. Bojja, M. Yaqub, F. Jinchao, and K. Arshid, \"Texture based localization of a brain tumor from MR-images by using a machine learning approach,\" Medical Hypotheses, vol. 141, p. 109705, 2020.",
    )
    add_text(doc, "Replace it with", bold=True)
    add_text(
        doc,
        "A. Mohebbi, S. Mohammadzadeh, A. H. Zare, Z. Moradi, A. Abbasian Ardakani, and A. Mohammadi, \"Assessing inter-rater reliability of MRI features in glioma: a multi-radiologist agreement study,\" BMC Medical Imaging, vol. 25, no. 1, art. no. 480, 2025, doi: 10.1186/s12880-025-01941-5.",
    )
    add_text(doc, "Delete this sentence", bold=True)
    add_text(
        doc,
        "Previous studies have reported high diagnostic agreement for general brain tumor assessment, but agreement may decrease in complex or mixed tumor cases [6].",
    )
    add_text(doc, "Replace it with", bold=True)
    add_text(
        doc,
        "Inter-rater agreement in brain tumor MRI interpretation can vary across imaging characteristics, with some glioma features showing high reliability and others, including diffusion characteristics, hemorrhage, and involvement of eloquent brain regions, showing poor agreement [6].",
    )

doc.save(OUTPUT)
print(OUTPUT)
