from pathlib import Path

from docx import Document


DOCX_PATH = Path(r"C:\Users\kaziNAHIAN\Downloads\notes.docx")
OUTPUT_PATH = Path(r"C:\Users\kaziNAHIAN\Downloads\notes_updated.docx")


def replace_paragraph(paragraph, text):
    for run in paragraph.runs:
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = text
    else:
        paragraph.add_run(text)


doc = Document(DOCX_PATH)

replacements = {
    168: (
        '[18] M. Musthafa, T. R. Mahesh, V. V. Kumar, and S. Guluwadi, '
        '"Enhancing brain tumor detection in MRI images through explainable AI using '
        'Grad-CAM with ResNet50," BMC Medical Imaging, vol. 24, Art. no. 107, 2024, '
        'doi: 10.1186/s12880-024-01292-7.'
    ),
    179: (
        '[47] A. M. Reza, "Realization of the contrast limited adaptive histogram '
        'equalization (CLAHE) for real-time image enhancement," The Journal of VLSI '
        'Signal Processing Systems for Signal, Image, and Video Technology, vol. 38, '
        'no. 1, pp. 35-44, 2004, doi: 10.1023/B:VLSI.0000028532.53893.82.'
    ),
    182: (
        'Replace the current NHS webpage with the following recent peer-reviewed review '
        'article from The Lancet. It provides an updated account of the heterogeneity, '
        'classification, diagnosis, and management of adult primary brain tumors.'
    ),
    185: (
        '[1] M. J. van den Bent, M. Geurts, P. J. French, M. Smits, D. Capper, '
        'J. E. C. Bromberg, and S. M. Chang, "Primary brain tumours in adults," '
        'The Lancet, vol. 402, no. 10412, pp. 1564-1579, 2023, '
        'doi: 10.1016/S0140-6736(23)01054-1.'
    ),
    189: (
        'Primary central nervous system tumors constitute a heterogeneous group of '
        'neoplasms with diverse molecular characteristics, biological behavior, and '
        'clinical outcomes [1].'
    ),
    196: (
        '[2] D. P. R. Chieffo, F. Lino, D. Ferrarese, D. Belella, G. M. Della Pepa, '
        'and F. Doglietto, "Brain tumor at diagnosis: From cognition and behavior to '
        'quality of life," Diagnostics, vol. 13, no. 3, Art. no. 541, 2023, '
        'doi: 10.3390/diagnostics13030541.'
    ),
    200: (
        "Brain tumors may cause neurological, cognitive, emotional, and behavioral "
        "disturbances that can substantially reduce patients' quality of life [2]."
    ),
}

for index, text in replacements.items():
    replace_paragraph(doc.paragraphs[index], text)

doc.save(OUTPUT_PATH)
print(OUTPUT_PATH)
